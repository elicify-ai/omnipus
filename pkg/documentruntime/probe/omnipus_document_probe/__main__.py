import argparse, hashlib, importlib, importlib.metadata, json, os, pathlib, shutil, subprocess, sys, tempfile, zipfile

FORMATS=("docx","xlsx","pptx","pdf")
IMPORTS={"docx":"docx","xlsx":"openpyxl","pptx":"pptx","pdf":"reportlab"}
DISTS={"docx":"python-docx","xlsx":"openpyxl","pptx":"python-pptx","pdf":"reportlab"}

def result(name, ok, detail): return {"format":name,"ok":ok,"detail":detail}
def fail(component, detail): print(json.dumps({"ok":False,"component":component,"detail":detail},ensure_ascii=False)); return 1

def check_version(actual, constraint):
    if not constraint: return True
    return actual == constraint.removeprefix("==")

def validate_asset(prefix, item):
    path=(prefix/item["path"]).resolve()
    if prefix not in path.parents: raise ValueError("asset escapes prefix")
    return hashlib.sha256(path.read_bytes()).hexdigest()==item["sha256"]

def generate(fmt, out):
    if fmt=="docx":
        from docx import Document
        d=Document(); d.add_heading("Résumé 東京",0); d.add_paragraph("Page one"); d.add_page_break(); d.add_paragraph("Page two"); d.save(out)
    elif fmt=="xlsx":
        from openpyxl import Workbook, load_workbook
        w=Workbook(); w.active["A1"]="Résumé 東京"; w.create_sheet("Deux")["A1"]="second"; w.save(out); load_workbook(out,read_only=True).close()
    elif fmt=="pptx":
        from pptx import Presentation
        p=Presentation(); s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text="Résumé 東京"; p.slides.add_slide(p.slide_layouts[6]); p.save(out)
    else:
        from reportlab.pdfgen.canvas import Canvas
        c=Canvas(str(out)); c.drawString(72,720,"Resume Tokyo"); c.showPage(); c.drawString(72,720,"Page two"); c.save()

def validate(fmt,path):
    data=path.read_bytes()
    if fmt=="pdf":
        from pypdf import PdfReader
        return data.startswith(b"%PDF-") and len(PdfReader(str(path)).pages)>=2
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist())
        required={"docx":"word/document.xml","xlsx":"xl/workbook.xml","pptx":"ppt/presentation.xml"}[fmt]
        return required in names and z.testzip() is None

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--manifest",required=True); ap.add_argument("--format",choices=("all",)+FORMATS,default="all"); ns=ap.parse_args()
    manifest_path=pathlib.Path(ns.manifest).resolve(); prefix=manifest_path.parent
    try: manifest=json.loads(manifest_path.read_text())
    except Exception as e: return fail("manifest",str(e))
    if manifest.get("revision") != prefix.name: return fail("manifest","revision/prefix mismatch")
    platform_name="windows" if sys.platform.startswith("win") else "darwin" if sys.platform=="darwin" else "linux" if sys.platform.startswith("linux") else sys.platform
    if platform_name not in manifest.get("platforms",[]): return fail("platform",f"unsupported {platform_name}")
    if prefix not in pathlib.Path(sys.executable).resolve().parents: return fail("python","interpreter resolves outside versioned prefix")
    for req in manifest.get("python_requirements",[]):
        try: importlib.import_module(req["import"]); actual=importlib.metadata.version(req.get("distribution",req["import"]))
        except Exception as e: return fail(req.get("import","python"),str(e))
        if not check_version(actual,req.get("version","")): return fail(req["import"],f"version {actual} does not satisfy {req['version']}")
    converter=pathlib.Path(manifest.get("converter",""))
    if not converter.is_file(): return fail("converter",f"missing {converter}")
    try: subprocess.run([str(converter),"--version"],check=True,capture_output=True,timeout=60)
    except Exception as e: return fail("converter",str(e))
    node=pathlib.Path(manifest.get("node",""))
    if not node.is_file(): return fail("node",f"missing {node}")
    try:
        node_path=subprocess.run([str(node),"-p","process.execPath"],check=True,capture_output=True,text=True,timeout=60).stdout.strip()
        if prefix not in pathlib.Path(node_path).resolve().parents: return fail("node","runtime resolves outside versioned prefix")
    except Exception as e: return fail("node",str(e))
    for asset in manifest.get("assets",[]):
        try:
            if not validate_asset(prefix,asset): return fail(asset["path"],"checksum mismatch")
        except Exception as e: return fail(asset.get("path","asset"),str(e))
    requested=FORMATS if ns.format=="all" else (ns.format,); results=[]
    with tempfile.TemporaryDirectory(prefix="omnipus-document-probe-",dir=os.getcwd()) as td:
        for fmt in requested:
            try:
                out=pathlib.Path(td)/("probe."+fmt); generate(fmt,out); ok=validate(fmt,out); results.append(result(fmt,ok,"generated and independently validated" if ok else "validation failed"))
            except Exception as e: results.append(result(fmt,False,str(e)))
    ok=all(r["ok"] for r in results); print(json.dumps({"ok":ok,"results":results},ensure_ascii=False)); return 0 if ok else 1

if __name__=="__main__": raise SystemExit(main())
