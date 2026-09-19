#!/usr/bin/env python3
"""Read-only formula/cache inventory. Not a calculation or fidelity validator."""
import argparse
import json
from pathlib import Path
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook


def inspect_workbook(path):
    path = Path(path)
    if path.suffix.lower() not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError('unsupported workbook extension')
    # Open first so permission and missing-file failures retain their own types.
    with path.open('rb') as source:
        try:
            with zipfile.ZipFile(source) as archive:
                names = set(archive.namelist())
            formulas = load_workbook(path, read_only=True, data_only=False, keep_links=True)
            try:
                cached = load_workbook(path, read_only=True, data_only=True, keep_links=True)
                try:
                    report = {'sheets': [], 'formulas': [], 'missing_formula_caches': [],
                              'errors': [], 'has_macros': 'xl/vbaProject.bin' in names,
                              'has_external_links': any(n.startswith('xl/externalLinks/externalLink') and n.endswith('.xml') for n in names)}
                    for sheet in formulas:
                        cache_sheet = cached[sheet.title]
                        # Producer-supplied bounds can be smaller than the stored cells.
                        sheet.reset_dimensions()
                        cache_sheet.reset_dimensions()
                        dimensions = {'name': sheet.title, 'rows': 1, 'columns': 1}
                        report['sheets'].append(dimensions)
                        for row, cache_row in zip(sheet.iter_rows(), cache_sheet.iter_rows()):
                            for cell, cache_cell in zip(row, cache_row):
                                if hasattr(cell, 'coordinate'):
                                    dimensions['rows'] = max(dimensions['rows'], cell.row)
                                    dimensions['columns'] = max(dimensions['columns'], cell.column)
                                location = {'sheet': sheet.title, 'cell': cell.coordinate} if cell.value is not None else None
                                if cell.data_type == 'f':
                                    report['formulas'].append(location)
                                    if cache_cell.value is None:
                                        report['missing_formula_caches'].append(location)
                                error = cell if cell.data_type == 'e' else cache_cell
                                if error.data_type == 'e':
                                    report['errors'].append({**location, 'value': error.value})
                    return report
                finally:
                    cached.close()
            finally:
                formulas.close()
        except (zipfile.BadZipFile, KeyError, ParseError, OSError) as exc:
            if isinstance(exc, OSError) and exc.errno is not None:
                raise
            raise ValueError('invalid workbook package') from exc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    args = parser.parse_args()
    try:
        print(json.dumps(inspect_workbook(args.workbook), indent=2))
    except (ValueError, OSError) as exc:
        parser.exit(2, f'Workbook inspection failed: {exc}\n')


if __name__ == '__main__':
    main()
