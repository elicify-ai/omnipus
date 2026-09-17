"""Read-only top-level shape bounds; not a rendered-overflow or overlap check."""
import argparse
import json
from zipfile import BadZipFile, ZipFile
from pptx import Presentation


def inspect_pptx(path):
    try:
        archive = ZipFile(path)
    except BadZipFile as exc:
        raise ValueError('not a ZIP-based presentation') from exc
    with archive:
        if 'ppt/presentation.xml' not in archive.namelist():
            raise ValueError('missing ppt/presentation.xml')
    deck = Presentation(path)
    width, height = deck.slide_width, deck.slide_height
    outside = []
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or
                    shape.left + shape.width > width or
                    shape.top + shape.height > height):
                outside.append({'slide': index, 'shape_id': shape.shape_id})
    return {'slide_count': len(deck.slides), 'width_emu': width,
            'height_emu': height, 'out_of_bounds': outside}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path')
    print(json.dumps(inspect_pptx(parser.parse_args().path), indent=2))
